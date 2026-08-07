"""
Export the Phase 3/4 result CSVs to Excel (.xlsx).

The thesis requirement is that every figure's underlying data is
available as an Excel file. This script collects all result CSVs in
results/ and writes them into a single workbook,
results/kudiloc_results.xlsx, with one sheet per result set. A single
workbook (rather than one .xlsx per CSV) keeps all the evidence for the
thesis in one place, easy to hand to a supervisor or attach as an
appendix, while still keeping each result set on its own tab.

Sheets written (when the corresponding CSV exists):
    adversarial_raw / adversarial_summary   (RQ2)
    staleness_raw   / staleness_summary     (RQ3)
    density_raw     / density_summary       (RQ1)
    halflife_sensitivity                    (robustness check)
    significance_tests                      (paired Wilcoxon results)

Formatting applied to each sheet:
    - Arial throughout (professional font, per output conventions).
    - Bold header row, frozen so it stays visible while scrolling.
    - Column widths sized to their contents.
    - Accuracy-style columns shown as percentages; p-values in
      scientific notation; small counts left as plain integers.

No formulas are written (this is a data export, not a model), so no
recalculation step is needed. Uses pandas for the bulk CSV->sheet load
and openpyxl for the light formatting pass.

Run from the repo root:
    PYTHONPATH=. python3 experiments/export_to_excel.py
"""

import os

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font

RESULTS_DIR = "results"
OUTPUT_PATH = os.path.join(RESULTS_DIR, "kudiloc_results.xlsx")

# (csv filename, sheet name). Order controls tab order in the workbook.
# Raw and summary are paired per research question so the tabs read in a
# sensible order for someone opening the file cold.
CSV_TO_SHEET = [
    ("adversarial_sweep.csv", "adversarial_raw"),
    ("adversarial_sweep_summary.csv", "adversarial_summary"),
    ("staleness_sweep.csv", "staleness_raw"),
    ("staleness_sweep_summary.csv", "staleness_summary"),
    ("density_sweep.csv", "density_raw"),
    ("density_sweep_summary.csv", "density_summary"),
    ("halflife_sensitivity.csv", "halflife_sensitivity"),
    ("significance_tests.csv", "significance_tests"),
]

# Columns whose values are proportions in [0, 1] and should display as
# percentages. Matched by exact column name across all sheets.
PERCENT_COLUMNS = {
    "accuracy",
    "mean_accuracy",
    "std_accuracy",
    "ci95_halfwidth",
    "mean_diff",
    "reports_per_atm",       # not a percent, but see override below
    "zero_report_fraction",
}
# reports_per_atm is a count, not a proportion, so exclude it from the
# percent set (kept above only to document the decision explicitly).
PERCENT_COLUMNS.discard("reports_per_atm")

# Columns to render in scientific notation (very small p-values).
SCIENTIFIC_COLUMNS = {"p_value"}

HEADER_FONT = Font(name="Arial", bold=True)
BODY_FONT = Font(name="Arial")


def write_sheets() -> list:
    """
    Load each existing CSV into a DataFrame and write it to its sheet.
    Returns the list of (sheet_name, dataframe) actually written, so the
    formatting pass knows which sheets and columns exist.
    """
    written = []

    with pd.ExcelWriter(OUTPUT_PATH, engine="openpyxl") as writer:
        for csv_name, sheet_name in CSV_TO_SHEET:
            csv_path = os.path.join(RESULTS_DIR, csv_name)
            if not os.path.exists(csv_path):
                print(f"SKIP: {csv_path} not found")
                continue

            df = pd.read_csv(csv_path)
            df.to_excel(writer, sheet_name=sheet_name, index=False)
            written.append((sheet_name, df))
            print(f"Wrote sheet '{sheet_name}' ({len(df)} rows) from {csv_name}")

    return written


def format_workbook(written: list) -> None:
    """
    Apply fonts, header styling, frozen header row, column widths, and
    number formats to every written sheet.
    """
    wb = load_workbook(OUTPUT_PATH)

    for sheet_name, df in written:
        ws = wb[sheet_name]

        # Map column name -> 1-based column index, from the header row.
        col_index = {col: i + 1 for i, col in enumerate(df.columns)}

        # Header row: bold Arial, and freeze it so it stays on screen.
        for cell in ws[1]:
            cell.font = HEADER_FONT
            cell.alignment = Alignment(horizontal="center")
        ws.freeze_panes = "A2"

        # Body: Arial everywhere, plus per-column number formats.
        for col_name, idx in col_index.items():
            for row in range(2, ws.max_row + 1):
                cell = ws.cell(row=row, column=idx)
                cell.font = BODY_FONT

                if cell.value is None or cell.value == "":
                    continue

                if col_name in PERCENT_COLUMNS:
                    cell.number_format = "0.00%"
                elif col_name in SCIENTIFIC_COLUMNS:
                    cell.number_format = "0.00E+00"

        # Column widths: fit the widest cell (header or body) in each
        # column, with a small margin, capped so a long text column does
        # not blow the layout out.
        for col_name, idx in col_index.items():
            max_len = len(str(col_name))
            for row in range(2, ws.max_row + 1):
                value = ws.cell(row=row, column=idx).value
                if value is not None:
                    max_len = max(max_len, len(str(value)))
            ws.column_dimensions[ws.cell(row=1, column=idx).column_letter].width = min(
                max_len + 2, 40
            )

    wb.save(OUTPUT_PATH)
    print(f"\nFormatted and saved {OUTPUT_PATH}")


def main():
    written = write_sheets()
    if not written:
        print("No CSVs found to export.")
        return
    format_workbook(written)


if __name__ == "__main__":
    main()
